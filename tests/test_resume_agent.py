import json

import pytest
from docx import Document

from app.resume_agent.claude_client import build_prompt, parse_edits
from app.resume_agent.docx_editor import apply_edits, extract_paragraphs


@pytest.fixture
def sample_docx(tmp_path):
    doc = Document()
    doc.add_paragraph("Objective: seeking a generic software role.")
    doc.add_paragraph("")  # blank paragraph should be skipped
    doc.add_paragraph("Built internal tools using Python and SQL.")
    path = tmp_path / "resume.docx"
    doc.save(path)
    return path


def test_extract_paragraphs_skips_blank(sample_docx):
    paragraphs = extract_paragraphs(str(sample_docx))
    assert len(paragraphs) == 2
    assert paragraphs[0] == {"index": 0, "text": "Objective: seeking a generic software role."}
    assert paragraphs[1] == {"index": 2, "text": "Built internal tools using Python and SQL."}


def test_apply_edits_overwrites_text_preserving_paragraph_count(sample_docx, tmp_path, monkeypatch):
    monkeypatch.setattr("app.resume_agent.docx_editor.TAILORED_DIR", tmp_path)
    edits = [{"index": 0, "new_text": "Objective: seeking a backend role using Python."}]

    output_path = apply_edits(str(sample_docx), edits, "test_output")

    result_doc = Document(output_path)
    assert len(result_doc.paragraphs) == 3
    assert result_doc.paragraphs[0].text == "Objective: seeking a backend role using Python."
    assert result_doc.paragraphs[2].text == "Built internal tools using Python and SQL."


def test_build_prompt_includes_job_description_and_paragraphs():
    paragraphs = [{"index": 0, "text": "Objective: foo"}]
    prompt = build_prompt(paragraphs, "We need a Python backend engineer.")
    assert "We need a Python backend engineer." in prompt
    assert "0: Objective: foo" in prompt
    assert "Never invent" in prompt


def test_parse_edits_valid_json():
    output = json.dumps({"summary": "Rewrote objective", "edits": [{"index": 0, "new_text": "New text"}]})
    summary, edits = parse_edits(output)
    assert summary == "Rewrote objective"
    assert edits == [{"index": 0, "new_text": "New text"}]


def test_parse_edits_strips_markdown_fence():
    output = '```json\n{"summary": "x", "edits": []}\n```'
    summary, edits = parse_edits(output)
    assert summary == "x"
    assert edits == []


def test_parse_edits_raises_on_invalid_json():
    with pytest.raises(ValueError):
        parse_edits("not json at all")


def test_parse_edits_raises_when_edits_missing():
    with pytest.raises(ValueError):
        parse_edits(json.dumps({"summary": "no edits key here"}))
